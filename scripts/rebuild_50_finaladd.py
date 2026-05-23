from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph
ROOT=Path.home()/"Desktop"
SRC=ROOT/"王昌振毕业论文v1.0_50页扩写版_正文未校页码3.docx"
OUT=ROOT/"王昌振毕业论文v1.0_50页扩写版_正文待校页码.docx"

def insert_after(p,text):
    new=deepcopy(p._p)
    for c in list(new): new.remove(c)
    p._p.addnext(new)
    np=Paragraph(new,p._parent); np.style='Normal'; np.add_run(text); return np

def body_start(doc):
    for i,p in enumerate(doc.paragraphs):
        if p.text.strip()=='4  系统详细设计与实现': return i
    raise ValueError

def find_body(doc,prefix):
    st=body_start(doc)
    for p in doc.paragraphs[st:]:
        if p.text.strip().startswith(prefix): return p
    raise ValueError(prefix)

def add(doc,prefix,arr):
    cur=find_body(doc,prefix)
    for t in arr: cur=insert_after(cur,t)

doc=Document(str(SRC))
add(doc,'5.3 局限',[
'从论文整体结构看，第 4 章和第 5 章是相互支撑的关系。第 4 章说明系统怎么设计、怎么实现，第 5 章说明这些实现如何被验证。扩写时我没有新增与项目无关的模块，而是围绕已经完成的设备管理、规则告警、AI 问答、工单流转、消息推送和预测服务继续展开，这样内容和项目实际代码能够对应。',
'后续如果继续完善系统，还可以把测试过程做得更加自动化。例如设备管理、工单状态流转和权限控制可以编写接口测试用例；AI 问答和预测服务可以保留固定输入样本，用于对比不同版本的返回结果。这样不仅方便论文验证，也方便项目后续迭代时快速发现回归问题。'
])
add(doc,'4.5 消息',[
'因此，本章最后可以概括为两条主线：一条是传统业务主线，包括设备、状态、规则、告警、工单和消息；另一条是智能辅助主线，包括知识库问答、自然语言查库和风电功率预测。两条主线通过统一的设备数据和用户权限连接起来，共同构成本文系统的主要实现内容。'
])
doc.save(str(OUT))
print(OUT)
print('chars',sum(len(p.text) for p in doc.paragraphs),'paragraphs',len(doc.paragraphs),'images',len(doc.inline_shapes),'tables',len(doc.tables))
