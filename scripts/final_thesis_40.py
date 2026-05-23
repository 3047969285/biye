from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph

DESKTOP=Path.home()/"Desktop"
SOURCE=DESKTOP/"王昌振毕业论文v1.0_设计实现分离修改版.docx"
OUTPUT=DESKTOP/"王昌振毕业论文v1.0_最终40页版.docx"

def insert_after(paragraph: Paragraph, text: str, style_name: str|None=None) -> Paragraph:
    new_p=deepcopy(paragraph._p)
    for child in list(new_p):
        new_p.remove(child)
    paragraph._p.addnext(new_p)
    new_para=Paragraph(new_p, paragraph._parent)
    if style_name:
        new_para.style=style_name
    new_para.add_run(text)
    return new_para

def set_text(p: Paragraph, text: str):
    for r in p.runs: r.text=''
    if p.runs: p.runs[0].text=text
    else: p.add_run(text)

def find_exact(doc, text):
    for p in doc.paragraphs:
        if p.text.strip()==text: return p
    raise ValueError(text)

def find_start(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix): return p
    raise ValueError(prefix)

def insert_chain(p, texts):
    cur=p
    for t in texts:
        cur=insert_after(cur,t,'Normal')
    return cur

def main():
    doc=Document(str(SOURCE))
    replacements={
        "4 系统详细实现\t14":"4 系统详细设计与实现\t14",
        "4  系统详细实现":"4  系统详细设计与实现",
        "上一章已经完成总体设计，本章不再重复设计内容，而是从代码落地角度说明系统实现。正文重点围绕 Controller、Service、Mapper 的职责划分，前端页面与后端接口的调用关系，以及 AI 问答、知识库检索、工单流转、风电预测和消息权限等功能在项目中的具体实现。":"本章将详细设计和具体实现分开说明。每个模块先说明设计思路、数据流向和边界控制，再说明后端接口、服务层、数据访问层及前端页面的实现方式，使系统从方案设计到代码落地的过程更加清晰。",
        "总体设计部分到此结束。图3.1展示系统分层关系，图3.2对应核心业务表之间的关联；第四章只展开实现细节，避免把架构设计、数据库设计和代码实现混在同一节中。":"总体设计部分到此结束。图3.1展示系统分层关系，图3.2对应核心业务表之间的关联；第四章在此基础上继续展开详细设计和实现，重点说明模块拆分、接口衔接和关键代码。"
    }
    for old,new in replacements.items():
        try: set_text(find_exact(doc,old),new)
        except ValueError: pass

    additions={
        "4.1":[
            "设计说明：设备数据采用台账信息和运行状态分离的方式组织。台账表保存设备名称、类型、位置等低频变化数据，状态表保存采集时间、指标值和状态描述等高频变化数据。二者通过设备编号关联，既便于维护基础资料，也便于追踪设备运行过程。",
            "实现说明：本模块按 Controller、Service、Mapper 三层落地。Controller 负责接收分页查询、新增、修改和删除请求，Service 负责参数校验、事务控制和缓存刷新，Mapper 负责具体 SQL。状态读取优先查缓存，未命中再回源数据库，保证查询效率和数据可靠性。"
        ],
        "4.2":[
            "设计说明：告警模块把状态写入和规则计算分开处理。状态数据进入系统后，只作为触发入口；规则服务再根据启用状态、阈值、持续条件和告警等级判断是否生成告警。这样后续增加新规则时，不需要改动设备状态接口。",
            "实现说明：触发服务先读取当前设备可用规则，再调用评估服务逐条判断。满足条件后写入告警记录，并调用消息服务生成站内提醒。对于短时间内重复出现的相同异常，系统会进行简单去重，减少运维人员看到重复告警的情况。"
        ],
        "4.3":[
            "设计说明：AI 模块分为自由对话、知识库问答和自然语言查库三种入口。自由对话用于解释运维问题，知识库问答用于结合上传资料回答，数据库问答用于把自然语言转换为只读查询。三种模式共用对话页面，但后端处理流程不同。",
            "实现说明：普通对话和 RAG 问答通过 Spring AI 调用模型，并用 SseEmitter 实现流式输出；知识库资料经过解析、切分、向量化后参与检索；数据库问答在执行前会校验 SQL 类型和危险关键词，避免模型生成修改或删除数据的语句。"
        ],
        "4.4":[
            "设计说明：运维表单用于承接告警后的处置流程，预测模块用于辅助提前研判。工单关注已经发生的问题如何分派、处理和归档，预测关注未来一段时间功率变化趋势。二者都围绕设备编号关联，便于在设备详情中形成完整记录。",
            "实现说明：工单服务限制状态跳转顺序，未分派工单不能直接完成，已归档工单不能继续修改。风电预测由 Java 后端调用 Python FastAPI 服务，Python 加载 GRU 模型完成推理，返回功率序列后由前端用 ECharts 绘制曲线。"
        ],
        "4.5":[
            "设计说明：消息、权限和日志是平台支撑能力。消息保证告警和工单变化能够及时触达，权限保证不同角色只能操作授权功能，日志保证关键操作可以追溯。它们不直接产生业务数据，但决定系统能否稳定、安全地使用。",
            "实现说明：告警和工单提醒采用站内消息并预留 WebSocket 推送能力，AI 对话采用 SSE 流式返回。权限控制同时放在前端按钮和后端接口上，前端通过权限指令控制显示，后端通过 @PreAuthorize 做最终拦截。"
        ]
    }
    for prefix,texts in additions.items():
        p=find_start(doc,prefix)
        insert_chain(p,texts)

    figure_notes={
        "图5.1.1":"图5.1.1对应设备修改测试，重点验证编辑窗口提交后，列表页和数据库记录是否同步更新，同时检查取消编辑时是否不会误写数据。",
        "图5.1.2":"图5.1.2对应设备新增测试，重点验证必填项校验、保存提示和新增后回显效果，说明设备主数据入口能够正常使用。",
        "图5.1.3":"图5.1.3对应基础对话测试，重点观察流式返回是否连续、页面是否卡顿以及对话历史是否能够保存。",
        "图5.1.4":"图5.1.4对应知识库问答测试，重点验证上传资料后系统能否检索相关片段，并围绕资料内容给出回答。",
        "图5.1.5":"图5.1.5对应数据库问答测试，重点验证自然语言生成只读查询后，后端能否完成安全校验并以表格方式展示结果。",
        "图5.1.6":"图5.1.6对应发电预测测试，重点验证 Java 后端能否调用 Python 预测服务，并把 GRU 模型返回的功率序列绘制成曲线。",
        "图5.1.7":"图5.1.7对应运维表单生成测试，重点验证工单创建、负责人填写、状态初始化和列表查询是否形成闭环。",
        "图5.1.8":"图5.1.8对应消息推送测试，重点验证告警或工单变化后，消息中心是否生成提醒并显示未读状态。",
        "图5.1.9":"图5.1.9对应告警管理测试，重点验证异常状态触发规则后，告警等级、设备信息和处理状态是否正确展示。"
    }
    for cap,note in figure_notes.items():
        insert_after(find_start(doc,cap),note,'Normal')

    # Replace a few generic test paragraphs with more detailed wording without adding too many pages.
    try:
        set_text(find_exact(doc,"整体来看，系统在当前环境下运行稳定，核心业务链路可用。"),
            "整体来看，系统在当前环境下运行稳定，核心业务链路可用。测试结果表明，设备数据、规则告警、AI 问答、工单处理、消息提醒和预测展示能够围绕设备运维形成连续流程。受限于硬件条件，本次测试仍以原型验证为主，后续接入真实采集网关后，还需要继续补充长时间运行和并发写入测试。")
    except ValueError:
        pass

    doc.save(str(OUTPUT))
    print(OUTPUT)
    print('chars',sum(len(p.text) for p in doc.paragraphs),'paragraphs',len(doc.paragraphs),'images',len(doc.inline_shapes))

if __name__=='__main__': main()
