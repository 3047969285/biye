from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph
ROOT=Path.home()/"Desktop"
SRC=ROOT/"王昌振毕业论文v1.0_50页扩写版_正文未校页码2.docx"
OUT=ROOT/"王昌振毕业论文v1.0_50页扩写版_正文未校页码3.docx"

def insert_after(p,text):
    new=deepcopy(p._p)
    for c in list(new): new.remove(c)
    p._p.addnext(new)
    np=Paragraph(new,p._parent); np.style='Normal'; np.add_run(text); return np

def body_start(doc):
    for i,p in enumerate(doc.paragraphs):
        if p.text.strip()=='4  系统详细设计与实现': return i
    raise ValueError

def find_body(doc,prefix):
    st=body_start(doc)
    for p in doc.paragraphs[st:]:
        if p.text.strip().startswith(prefix): return p
    raise ValueError(prefix)

def add(doc,prefix,arr):
    cur=find_body(doc,prefix)
    for t in arr: cur=insert_after(cur,t)

doc=Document(str(SRC))
add(doc,'4.3 对话',[
'知识库问答在实现中还要考虑资料更新问题。毕业设计阶段主要是上传资料后立即使用，但如果后续资料经常变化，就需要重新解析文档并更新向量片段。为此，文档记录表中保存了文件名、上传时间和分类等信息，后续可以根据这些信息判断哪些资料需要重新入库。这个设计为后续扩展留下了空间。',
'自然语言查库的结果展示也需要尽量简单。运维人员关心的是查询结果本身，而不是 SQL 语句有多复杂。因此前端把查询结果按表格展示，并保留必要说明。对于没有结果的查询，系统返回空结果提示；对于不允许执行的语句，系统给出安全限制提示。这样用户能够理解系统反馈，而不是看到一段难以理解的报错。',
'AI 模块和传统业务模块结合时，最重要的是边界清楚。模型可以帮助解释、检索和生成查询，但不能绕过权限，也不能直接修改核心数据。本系统把 AI 能力放在受控接口之后，所有数据库访问仍然经过后端校验，这样既能体现智能化特点，又能保留管理系统应有的安全边界。'
])
add(doc,'4.4 运维',[
'工单详情页在设计上不仅展示当前状态，也需要展示问题来源和处理过程。这样运维人员接手时可以知道工单是人工创建还是由告警转化而来，之前已经填写过哪些处理意见，当前还需要做什么。这个信息组织方式对多人协作比较重要，也方便后续追溯。',
'预测模块在页面展示时，曲线比表格更容易表达趋势。功率数据本身是一组连续数值，如果只放在表格里，用户很难看出变化方向；使用 ECharts 曲线后，可以直观看到上升、下降和波动区间。对于风电场景来说，趋势信息比单个数值更有参考意义。',
'预测服务和工单模块结合后，可以形成“事前研判”和“事后处理”的互补关系。预测结果用于提示未来可能的变化，工单记录用于保存已经发生问题的处理过程。虽然当前系统还没有实现自动根据预测生成工单，但通过统一设备编号，已经具备继续扩展的基础。'
])
add(doc,'5.2 性能',[
'在页面响应方面，设备列表和工单列表的操作比较稳定，主要原因是分页查询减少了一次性加载的数据量。AI 问答和预测服务的耗时相对更不稳定，因为它们依赖外部模型接口和 Python 推理服务。测试时我把这两类功能分开观察，避免用同一个标准评价所有模块。',
'对于异常场景，测试并不是为了故意制造错误，而是为了确认系统在出错时能给出可理解的反馈。比如预测服务没有启动时，用户需要知道是预测服务问题，而不是误以为设备数据错误；数据库问答被拦截时，也需要说明是安全限制，而不是系统崩溃。'
])
doc.save(str(OUT))
print(OUT)
print('chars',sum(len(p.text) for p in doc.paragraphs),'paragraphs',len(doc.paragraphs),'images',len(doc.inline_shapes),'tables',len(doc.tables))
