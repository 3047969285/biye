from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph
ROOT=Path.home()/"Desktop"
SRC=ROOT/"王昌振毕业论文v1.0_50页扩写版_待校页码4.docx"
OUT=ROOT/"王昌振毕业论文v1.0_50页扩写版_待校页码5.docx"

def insert_after(p,text):
    new=deepcopy(p._p)
    for c in list(new): new.remove(c)
    p._p.addnext(new)
    np=Paragraph(new,p._parent); np.style='Normal'; np.add_run(text); return np

def find_start(doc,prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix): return p
    raise ValueError(prefix)

def add(p,arr):
    cur=p
    for t in arr: cur=insert_after(cur,t)

doc=Document(str(SRC))
add(find_start(doc,"4.5 消息"),[
"从整体实现看，第 4 章中的几个模块并不是孤立存在的。设备台账提供基础对象，状态数据提供运行依据，规则服务把状态变化转化为告警，告警继续推动消息和工单，AI 问答用于辅助理解和查询，预测模块用于补充趋势判断。这样的组织方式使系统具有一条比较完整的业务主线，而不是简单把若干页面放在同一个后台模板中。",
"在开发过程中，前后端字段命名的一致性也很重要。比如设备编号、工单状态、预测结果数组等字段，如果前端使用的名称和后端返回不一致，页面就会出现空白或图表无法渲染。为了解决这类问题，我在联调时会先用接口工具查看返回 JSON，再对照前端字段绑定逐项检查。这个过程虽然比较细，但能有效减少页面显示异常。",
"另外，系统中涉及外部能力的部分都尽量增加了容错说明。AI 模型接口、知识库检索和 Python 预测服务都可能受到网络、配置或文件路径影响。后端在这些地方返回可读提示，前端再用消息框展示出来，用户能够知道是服务不可用、参数错误还是没有查询结果。这样的处理比直接报错更适合管理系统。",
"本章加入代码片段的目的不是展示全部源码，而是选取最能体现项目特点的位置。设备状态写入体现事务和缓存，规则触发体现告警联动，SSE 代码体现流式问答，预测接口体现 Java 与 Python 的连接，权限注解体现前后端权限控制。通过这些代码，读者能够看到系统关键功能确实有对应实现。"
])
doc.save(str(OUT))
print(OUT)
print('chars',sum(len(p.text) for p in doc.paragraphs),'paragraphs',len(doc.paragraphs),'images',len(doc.inline_shapes),'tables',len(doc.tables))
