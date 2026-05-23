from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph
ROOT=Path.home()/"Desktop"
SRC=ROOT/"王昌振毕业论文v1.0_50页扩写版_二次扩写.docx"
OUT=ROOT/"王昌振毕业论文v1.0_50页扩写版_待校页码.docx"

def insert_after(p,text):
    new=deepcopy(p._p)
    for c in list(new): new.remove(c)
    p._p.addnext(new)
    np=Paragraph(new,p._parent); np.style='Normal'; np.add_run(text); return np

def find_start(doc,prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix): return p
    raise ValueError(prefix)

def add(p,arr):
    cur=p
    for t in arr: cur=insert_after(cur,t)

doc=Document(str(SRC))
add(find_start(doc,"4.3 对话"),[
"在自然语言查库测试中，我没有把它设计成万能数据库助手，而是只围绕本系统已有业务表进行查询。这样做的原因是毕业设计系统数据范围有限，过度开放反而会带来安全风险。通过限制查询范围和返回数量，系统既能体现 AI 查询的便利，也能保持数据库访问相对可控。",
"对话历史的保存也有实际意义。运维场景中，一个问题往往不是一次提问就能完全解决，用户可能需要结合设备信息、知识库资料和历史处理记录多次追问。保存历史记录后，用户重新进入页面还能回看之前的问答内容，这比单纯调用模型接口更接近完整应用。"
])
add(find_start(doc,"4.4 运维"),[
"工单和预测虽然属于不同功能，但在系统中都围绕设备展开。工单记录过去和当前的问题处理情况，预测结果提供未来一段时间的趋势参考。把二者都关联到设备详情中，用户就可以在同一个入口看到设备基础信息、历史异常、处理记录和预测曲线，信息链条更加完整。",
"预测结果保存到数据库后，也为后续对比分析留下了空间。当前论文主要展示本次预测曲线，后续可以继续把历史预测值和真实运行值放在一起比较，计算不同时间段的误差变化。这样预测模块就不只是一次性演示，而可以逐步形成模型效果评估数据。"
])
add(find_start(doc,"5.1 测试"),[
"测试章节扩写时，我尽量把截图和文字说明对应起来，避免只放图不解释。每张图都说明了测试入口、观察重点和能够证明的功能点。这样阅读论文时，即使不运行项目，也能从截图和说明中看出系统主要功能是如何完成的。",
"在整理测试结果时，我把“能打开页面”和“业务流程可用”区分开来。页面能打开只能说明前端路由没有问题，业务流程可用还需要验证接口返回、数据库保存、状态变化和前端回显。本文第 5 章更关注后者，因此对工单、告警、预测和 AI 问答都补充了流程性说明。"
])
doc.save(str(OUT))
print(OUT)
print('chars',sum(len(p.text) for p in doc.paragraphs),'paragraphs',len(doc.paragraphs),'images',len(doc.inline_shapes),'tables',len(doc.tables))
