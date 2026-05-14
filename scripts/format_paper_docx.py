from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path("C:/Users/Administrator/Desktop/毕业/毕业答辩/王昌振毕业论文（1） - Codex修订版.docx")
TARGET = Path("C:/Users/Administrator/Desktop/毕业/毕业答辩/王昌振毕业论文（1） - Codex修订目录三线表版.docx")


H1 = {
    "1  绪论",
    "2  系统需求分析",
    "3  系统总体设计",
    "4  系统详细设计与实现",
    "5  系统测试",
    "6  总结与展望",
    "参考文献",
    "致  谢",
}
H2_PREFIXES = ("1.1", "1.2", "1.3", "1.4", "2.1", "2.2", "2.3", "3.1", "3.2", "4.1", "4.2", "4.3", "4.4", "4.5", "5.1", "5.2", "5.3")
H3_RE = re.compile(r"^\d+\.\d+\.\d+\s+")


def insert_paragraph_before(paragraph, text, style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p = new_p
    new_para._element = new_p
    if style:
        new_para.style = style
    new_para.add_run(text)
    return new_para


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def set_run_font(paragraph, size=Pt(12), bold=None, name="宋体"):
    for run in paragraph.runs:
        run.font.name = name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
        run.font.size = size
        if bold is not None:
            run.bold = bold


def apply_heading_style(paragraph, level):
    paragraph.style = f"Heading {level}"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.5
    set_run_font(paragraph, Pt(16 if level == 1 else 14 if level == 2 else 12), True, "黑体")


def normalize_body_style(paragraph):
    if paragraph.text.strip() and paragraph.style.name in {"论文1", "论文2", "Heading 1", "Heading 2", "Heading 3", "toc 1", "toc 2", "toc 3"}:
        paragraph.style = "Normal"
        paragraph.paragraph_format.first_line_indent = Pt(24)
        paragraph.paragraph_format.line_spacing = 1.5
        set_run_font(paragraph, Pt(12), False, "宋体")


def clear_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, attrs in kwargs.items():
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in attrs.items():
            element.set(qn(f"w:{key}"), str(value))


def apply_three_line_table(table):
    clear_table_borders(table)
    rows = table.rows
    if not rows:
        return
    line = {"val": "single", "sz": "8", "space": "0", "color": "000000"}
    thin = {"val": "single", "sz": "6", "space": "0", "color": "000000"}
    for row in rows:
        for cell in row.cells:
            set_cell_border(
                cell,
                top={"val": "nil"},
                left={"val": "nil"},
                bottom={"val": "nil"},
                right={"val": "nil"},
                insideH={"val": "nil"},
                insideV={"val": "nil"},
            )
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_run_font(p, Pt(10.5), None, "宋体")
    for cell in rows[0].cells:
        set_cell_border(cell, top=line, bottom=thin)
        for p in cell.paragraphs:
            set_run_font(p, Pt(10.5), True, "宋体")
    for cell in rows[-1].cells:
        set_cell_border(cell, bottom=line)


def fix_content(doc):
    # 删除旧的手工目录，只保留“目  录”标题，后续由 Word 自动生成目录。
    paragraphs = list(doc.paragraphs)
    toc_title = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "目  录")
    body_start = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "1  绪论")
    for paragraph in paragraphs[toc_title + 1 : body_start]:
        remove_paragraph(paragraph)

    # 重新读取段落，修补缺失标题和误拼接标题。
    paragraphs = list(doc.paragraphs)
    p187 = next((p for p in paragraphs if "4.2 规则评估与告警联动" in p.text), None)
    if p187 is not None:
        p187.text = p187.text.replace("4.2 规则评估与告警联动", "").rstrip()
        target = next(
            p
            for p in doc.paragraphs
            if p.text.strip().startswith("规则评估与告警联动模块")
        )
        insert_paragraph_before(target, "4.2 规则评估与告警联动")

    if not any(p.text.strip() == "1.3 研究内容与研究意义" for p in doc.paragraphs):
        insert_paragraph_before(next(p for p in doc.paragraphs if p.text.strip() == "1.3.1 研究内容"), "1.3 研究内容与研究意义")
    if not any(p.text.strip() == "5.1 测试环境与功能验证" for p in doc.paragraphs):
        insert_paragraph_before(next(p for p in doc.paragraphs if p.text.strip().startswith("本章对系统各功能模块进行了测试验证")), "5.1 测试环境与功能验证")


def format_headings_and_body(doc):
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text or text == "目  录":
            continue
        normalize_body_style(paragraph)
        if text in H1:
            apply_heading_style(paragraph, 1)
        elif H3_RE.match(text):
            apply_heading_style(paragraph, 3)
        elif text.startswith(H2_PREFIXES):
            apply_heading_style(paragraph, 2)


def update_toc_with_word(path):
    import win32com.client

    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    doc = word.Documents.Open(str(path))
    try:
        for toc in list(doc.TablesOfContents):
            toc.Delete()
        toc_title = None
        for paragraph in doc.Paragraphs:
            if paragraph.Range.Text.strip() == "目  录":
                toc_title = paragraph
                break
        if toc_title is None:
            raise RuntimeError("未找到目录标题")
        rng = toc_title.Range
        rng.Collapse(0)  # wdCollapseEnd
        doc.TablesOfContents.Add(Range=rng, UseHeadingStyles=True, UpperHeadingLevel=1, LowerHeadingLevel=3, IncludePageNumbers=True, RightAlignPageNumbers=True)
        doc.TablesOfContents(1).Update()
        doc.Fields.Update()
        doc.Save()
    finally:
        doc.Close(SaveChanges=True)
        word.Quit()


def main():
    doc = Document(str(SOURCE))
    fix_content(doc)
    format_headings_and_body(doc)
    for index, table in enumerate(doc.tables):
        if index == 0:
            continue
        apply_three_line_table(table)
    doc.save(str(TARGET))
    update_toc_with_word(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
