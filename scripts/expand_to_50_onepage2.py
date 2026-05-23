from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph
ROOT=Path.home()/"Desktop"
SRC=ROOT/"王昌振毕业论文v1.0_50页扩写版_待校页码2.docx"
OUT=ROOT/"王昌振毕业论文v1.0_50页扩写版_待校页码3.docx"

def insert_after(p,text):
    new=deepcopy(p._p)
    for c in list(new): new.remove(c)
    p._p.addnext(new)
    np=Paragraph(new,p._parent); np.style='Normal'; np.add_run(text); return np

def find_start(doc,prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix): return p
    raise ValueError(prefix)

doc=Document(str(SRC))
p=find_start(doc,"5.3 局限")
insert_after(p,"从论文写作角度看，第 5 章的测试说明主要用于支撑前文实现内容。第四章写到的设备缓存、规则触发、SSE 对话、RAG 检索、数据库问答、工单流转和预测服务，在第五章都能找到对应测试入口和截图说明。这样章节之间能够相互呼应，避免实现部分写得较多而测试部分只有简单截图。")
doc.save(str(OUT))
print(OUT)
print('chars',sum(len(p.text) for p in doc.paragraphs),'paragraphs',len(doc.paragraphs),'images',len(doc.inline_shapes),'tables',len(doc.tables))
