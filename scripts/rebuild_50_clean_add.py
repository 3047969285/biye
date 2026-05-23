from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph
ROOT=Path.home()/"Desktop"
SRC=ROOT/"王昌振毕业论文v1.0_50页扩写版_正文未校页码.docx"
OUT=ROOT/"王昌振毕业论文v1.0_50页扩写版_正文未校页码2.docx"

def insert_after(p,text):
    new=deepcopy(p._p)
    for c in list(new): new.remove(c)
    p._p.addnext(new)
    np=Paragraph(new,p._parent); np.style='Normal'; np.add_run(text); return np

def body_start(doc):
    for i,p in enumerate(doc.paragraphs):
        if p.text.strip()=='4  系统详细设计与实现': return i
    raise ValueError

def find_body(doc,prefix):
    st=body_start(doc)
    for p in doc.paragraphs[st:]:
        if p.text.strip().startswith(prefix): return p
    raise ValueError(prefix)

def add(doc,prefix,arr):
    cur=find_body(doc,prefix)
    for t in arr: cur=insert_after(cur,t)

doc=Document(str(SRC))
add(doc,'4.5 消息',[
'在前后端联调时，权限标识的命名也需要保持稳定。如果后端接口使用 equipment:device:edit，而前端按钮写成另一个标识，就会出现按钮显示和接口权限不一致的问题。因此本系统尽量沿用“模块:功能:操作”的命名方式，方便在菜单、按钮和接口之间建立对应关系。这个细节虽然不影响单个功能演示，但会影响系统后期维护。',
'消息模块还承担了用户操作反馈的一部分作用。告警、工单和系统通知都可以通过消息中心统一查看，用户不需要在多个页面之间反复寻找待办内容。对于运维类系统来说，及时提醒比单纯保存记录更重要，因为异常处理往往具有时间要求，提醒越清晰，问题越不容易被遗漏。',
'从工程角度看，第 4 章的实现内容体现了多模块协同。后端有 Spring Boot、MyBatis、Redis、Spring Security 和 Spring AI，前端有 Vue、Element UI、Axios 和 ECharts，预测部分还接入了 Python 服务。论文中把这些内容放在不同小节，是为了说明每个技术点对应的业务位置，而不是简单罗列技术栈。'
])
add(doc,'5.1 测试',[
'测试截图整理完成后，我又对正文中的图号进行了逐一核对，确保每一张测试图在正文中都有对应说明。这样做可以避免论文中出现“只有图、没有解释”的情况。尤其是第 5 章的图比较多，如果不说明每张图验证了什么，读者很难判断这些截图和系统功能之间的关系。',
'在测试记录中，我更关注操作前后的变化。例如设备修改要看修改前字段、保存提示和修改后回显；工单测试要看状态变化；预测测试要看请求后是否生成曲线；消息测试要看未读数量和消息内容是否变化。通过这些前后对比，测试结论会比单纯展示页面截图更可靠。',
'测试章节还起到承接第 4 章实现内容的作用。第四章写了代码和实现方式，第五章就需要用测试结果说明这些实现能够运行。因此本文把设备、告警、AI、数据库问答、预测、工单、消息和权限都放入测试范围，尽量覆盖项目中比较重要的功能点。'
])
doc.save(str(OUT))
print(OUT)
print('chars',sum(len(p.text) for p in doc.paragraphs),'paragraphs',len(doc.paragraphs),'images',len(doc.inline_shapes),'tables',len(doc.tables))
