from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph
ROOT=Path.home()/"Desktop"
SRC=ROOT/"王昌振毕业论文v1.0_50页扩写版_待校页码.docx"
OUT=ROOT/"王昌振毕业论文v1.0_50页扩写版_待校页码2.docx"

def insert_after(p,text):
    new=deepcopy(p._p)
    for c in list(new): new.remove(c)
    p._p.addnext(new)
    np=Paragraph(new,p._parent); np.style='Normal'; np.add_run(text); return np

def find_start(doc,prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix): return p
    raise ValueError(prefix)

doc=Document(str(SRC))
p=find_start(doc,"5.2 性能")
insert_after(p,"综合来看，本次测试更强调项目在本机环境下的完整运行效果。设备、告警、AI、工单、消息和预测功能虽然复杂度不同，但都需要经过前端操作、后端接口、数据库或外部服务返回这几个环节。测试过程中没有把单个截图孤立看待，而是结合操作步骤和返回结果判断功能是否真正完成，这也是本文测试章节补充文字说明的主要原因。")
doc.save(str(OUT))
print(OUT)
print('chars',sum(len(p.text) for p in doc.paragraphs),'paragraphs',len(doc.paragraphs),'images',len(doc.inline_shapes),'tables',len(doc.tables))
