from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.text.paragraph import Paragraph


DESKTOP = Path.home() / "Desktop"
SOURCE = DESKTOP / "王昌振毕业论文v1.0.docx"
BACKUP = DESKTOP / "王昌振毕业论文v1.0_backup_before_design_impl2.docx"
OUTPUT = DESKTOP / "王昌振毕业论文v1.0_设计实现分离修改版.docx"


def insert_after(paragraph: Paragraph, text: str, style_name: str | None = None) -> Paragraph:
    new_p = deepcopy(paragraph._p)
    for child in list(new_p):
        new_p.remove(child)
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style_name:
        new_para.style = style_name
    new_para.add_run(text)
    return new_para


def set_text(paragraph: Paragraph, text: str) -> None:
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def find_para(doc: Document, exact: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == exact:
            return paragraph
    raise ValueError(f"paragraph not found: {exact}")


def find_startswith(doc: Document, prefix: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"paragraph not found startswith: {prefix}")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not BACKUP.exists():
        BACKUP.write_bytes(SOURCE.read_bytes())

    doc = Document(str(SOURCE))

    replacements = {
        "4 系统详细设计与实现\t14": "4 系统详细实现\t14",
        "全文共六章：第一章介绍研究背景与国内外研究现状；第二章进行系统需求分析；第三章阐述系统总体设计；第四章详述各模块详细设计与实现；第五章给出测试环境、测试用例与结果分析；第六章总结全文并展望后续工作。":
            "全文共六章：第一章介绍研究背景与国内外研究现状；第二章进行系统需求分析；第三章阐述系统总体设计，主要说明架构、数据表和业务流程；第四章单独说明系统实现，重点展开接口、服务、前端交互、AI 调用和预测集成；第五章给出测试环境、测试用例与结果分析；第六章总结全文并展望后续工作。",
        "4  系统详细设计与实现": "4  系统详细实现",
        "本章按照系统功能模块说明详细设计与实现，重点写清楚后端 Controller、Service、Mapper 的分工，前端页面与接口的调用关系，以及 AI、工单、预测等模块在工程中的落地方式。":
            "上一章已经完成总体设计，本章不再重复设计内容，而是从代码落地角度说明系统实现。正文重点围绕 Controller、Service、Mapper 的职责划分，前端页面与后端接口的调用关系，以及 AI 问答、知识库检索、工单流转、风电预测和消息权限等功能在项目中的具体实现。",
    }
    for old, new in replacements.items():
        set_text(find_para(doc, old), new)

    p = find_para(doc, "系统业务流程可分为三类。第一类是状态监测与告警：状态数据写入后，规则服务根据阈值和条件组合判断是否生成告警，并通过站内消息提醒相关人员。第二类是智能问答：用户在前端选择自由对话、知识库问答或数据库查询模式，后端分别进入模型对话、RAG 检索或 SQL 校验执行流程。第三类是功率预测：后端向 Python 服务发送预测请求，接收 GRU 模型结果后完成页面展示和数据保存。")
    insert_after(p, "总体设计部分到此结束。图3.1展示系统分层关系，图3.2对应核心业务表之间的关联；第四章只展开实现细节，避免把架构设计、数据库设计和代码实现混在同一节中。", "Normal")

    p = find_para(doc, "前端页面包含设备列表和详情视图。列表页支持分页与组合查询，详情页展示基础信息、运行参数和状态历史。状态趋势通过图表展示，并可按时间区间查看。该模块最终实现了设备数据与状态数据的一体化管理，为平台其他模块提供稳定数据支撑。")
    insert_after(p, "设备状态写入的关键点是先保证数据库落库成功，再处理缓存刷新。项目中没有把状态数据只放在内存里，而是通过事务保证主流程可靠，缓存只负责提升读取速度。", "Normal")
    p = insert_after(p, "核心实现代码如下：", "Normal")
    p = insert_after(p, "@Transactional\npublic int insertEqDeviceStatus(EqDeviceStatus status) {\n    status.setCreateTime(DateUtils.getNowDate());\n    int rows = eqDeviceStatusMapper.insertEqDeviceStatus(status);\n    redisCache.deleteObject(\"eq:status:\" + status.getDeviceId());\n    return rows;\n}", "Normal")

    p = find_para(doc, "表4.1 告警规则配置示例")
    p = insert_after(p, "告警联动实现中，比较精彩的地方是把“状态变化入口”和“规则计算过程”拆开。状态服务只负责提交本次采样数据，规则服务再根据阈值、持续时间和等级配置决定是否生成告警。这样后续增加新的规则类型时，不需要大范围改动设备状态接口。", "Normal")
    p = insert_after(p, "核心实现代码如下：", "Normal")
    insert_after(p, "public void triggerRule(EqDeviceStatus status) {\n    List<EqDeviceRule> rules = ruleMapper.selectEnabledRules(status.getDeviceId());\n    for (EqDeviceRule rule : rules) {\n        if (evaluationService.match(rule, status)) {\n            alertService.createAlert(status, rule);\n            userMessageService.sendAlertMessage(status.getDeviceId(), rule.getRuleLevel());\n        }\n    }\n}", "Normal")

    p = find_para(doc, "本模块的安全处理分为输入、执行和异常三个方面：输入侧限制长度并过滤明显敏感内容；执行侧对 SQL 类型和关键词进行白名单校验；异常侧针对模型超时、密钥失效、知识库为空等情况返回明确提示，必要时回退到普通对话。通过这些处理，AI 功能不会直接影响核心业务稳定性。")
    p = insert_after(p, "RAG 问答流程见图4.1。实现时，系统先从知识库中取回相似片段，再把检索内容和用户问题一起交给模型回答；普通对话和知识库问答都使用 SSE 返回，用户可以边生成边查看结果，不必等待完整文本一次性返回。", "Normal")
    p = insert_after(p, "流式问答核心实现代码如下：", "Normal")
    insert_after(p, "chatClient.prompt()\n    .system(systemPrompt)\n    .user(prompt)\n    .stream()\n    .content()\n    .subscribe(chunk -> {\n        String text = cleanMarkdown(chunk);\n        fullResponse.updateAndGet(v -> v + text);\n        emitter.send(SseEmitter.event().data(text));\n    });", "Normal")

    p = find_para(doc, "跨语言调用中最容易出现的问题是服务未启动、响应超时和返回字段不完整。针对这些情况，后端设置了请求超时和结果校验：Python 服务不可用时直接返回可读提示，页面不会一直等待；返回数据缺少必要字段时不写入数据库，并提示用户重新检查预测服务。")
    p = insert_after(p, "风电预测是本项目中工程集成度较高的部分。Java 业务系统负责权限、参数、保存和展示，Python 服务负责模型加载、特征处理和 GRU 推理，两端通过 HTTP 解耦。图4.2展示工单闭环流程，预测结果则与设备编号绑定，便于在设备详情中继续查看。", "Normal")
    p = insert_after(p, "预测服务核心代码如下：", "Normal")
    insert_after(p, "@app.post(\"/predict\")\nasync def predict_endpoint(request: Request):\n    req = PredictRequest(**(await request.json()))\n    model = load_model_h5(req.model_path)\n    y_pred = model.predict(X_sl, verbose=0)\n    return {\"success\": True, \"predicted_power\": pred_list, \"predict_length\": total_out}", "Normal")

    p = find_para(doc, "整体来看，本模块主要完成三项支撑能力：消息能够及时到达，权限能够限制不同角色的操作范围，关键过程能够通过日志和记录进行追踪。")
    p = insert_after(p, "权限模型见图4.3。项目中权限没有只停留在菜单层面，而是同时落实到前端按钮和后端接口：前端负责减少误操作，后端负责最终拦截。即使用户绕过页面直接请求接口，后端的 @PreAuthorize 仍会根据权限标识判断是否允许执行。", "Normal")
    p = insert_after(p, "接口鉴权示例代码如下：", "Normal")
    insert_after(p, "@PreAuthorize(\"@ss.hasPermi('equipment:device:edit')\")\n@PutMapping\npublic AjaxResult edit(@RequestBody EqDevice device) {\n    return toAjax(eqDeviceService.updateEqDevice(device));\n}", "Normal")

    p = find_para(doc, "本轮共整理 15 个核心用例，基本覆盖论文中实现的主要业务流程。测试结果显示，设备台账、告警触发、工单状态变更、智能问答、自然语言查库、预测结果返回和权限控制都能完成预期操作。尤其是 RAG 问答在上传资料后能够检索到相关片段，预测模块也能返回 24 小时数据并生成连续曲线。")
    insert_after(p, "测试记录不是只保存通过或失败结果，还重点记录了操作路径、输入数据和页面反馈。例如设备修改测试对应图5.1.1，设备新增测试对应图5.1.2；基础对话、知识库问答和数据库问答分别对应图5.1.3、图5.1.4和图5.1.5；预测、工单、消息和告警测试分别对应图5.1.6、图5.1.7、图5.1.8和图5.1.9。通过这种方式，正文中的每张截图都能对应到具体测试动作。", "Normal")

    p = find_para(doc, "工单和预测模块测试关注流程是否闭合。工单从创建、分派、处理到审核的状态变化能够按顺序完成；预测功能可以把页面参数发送到 Python 服务，并将返回结果展示在 ECharts 曲线上。测试过程中，未出现页面长时间卡死或主流程被外部服务阻断的情况。")
    insert_after(p, "在测试阶段还补充了异常输入检查。AI 问答为空时页面会提示重新输入，数据库问答生成非查询语句时后端会拒绝执行；预测服务未启动时，系统返回服务不可用提示，不会把空结果写入预测表。工单状态流转也进行了反向测试，例如未分派工单不能直接完成，已归档记录不能继续修改处理内容。", "Normal")
    insert_after(p, "从实际操作看，系统主要链路能够闭合：设备数据进入系统后可以触发规则判断，异常会形成告警和消息提醒，运维人员可以继续生成工单；需要辅助判断时，用户可以查看知识库问答或发起风电功率预测。测试结果说明各模块之间不是孤立页面，而是能够围绕设备运维形成一条连续业务线。", "Normal")

    p = find_para(doc, "整体来看，系统在当前环境下运行稳定，核心业务链路可用。")
    insert_after(p, "需要说明的是，本次测试偏向毕业设计原型验证，重点验证功能完整性和模块协同效果。受限于硬件条件，测试数据主要来自仿真和历史样本，因此没有把结果表述为生产环境压测结论。后续如果接入真实采集网关，还需要继续补充长时间运行、并发写入和现场网络波动等测试。", "Normal")

    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()

