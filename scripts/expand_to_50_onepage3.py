from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph
ROOT=Path.home()/"Desktop"
SRC=ROOT/"王昌振毕业论文v1.0_50页扩写版_待校页码3.docx"
OUT=ROOT/"王昌振毕业论文v1.0_50页扩写版_待校页码4.docx"

def insert_after(p,text):
    new=deepcopy(p._p)
    for c in list(new): new.remove(c)
    p._p.addnext(new)
    np=Paragraph(new,p._parent); np.style='Normal'; np.add_run(text); return np

def find_start(doc,prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix): return p
    raise ValueError(prefix)

doc=Document(str(SRC))
p=find_start(doc,"5.2 性能")
insert_after(p,"在功能复测时，我还对不同模块之间的跳转关系进行了检查。例如从设备信息进入相关状态记录，从告警记录继续查看处理情况，从工单列表进入详情页，再返回列表观察状态是否刷新。虽然这些操作不属于单个算法或接口，但它们直接影响系统是否能连续使用。测试结果表明，主要页面之间的跳转和刷新逻辑能够满足毕业设计展示需要，没有出现保存成功但页面仍显示旧数据的明显问题。")
insert_after(p,"测试过程也暴露出一些需要注意的地方。AI 和预测都依赖外部服务，调试时必须确认密钥、网络、Python 环境和模型文件路径是否正确；如果只看前端页面，很难快速判断问题来源。因此在测试中我同时观察浏览器请求、后端日志和 Python 服务输出，把问题定位到具体环节。这部分经历也说明，本项目不是单纯页面拼接，而是涉及多服务协同。")
doc.save(str(OUT))
print(OUT)
print('chars',sum(len(p.text) for p in doc.paragraphs),'paragraphs',len(doc.paragraphs),'images',len(doc.inline_shapes),'tables',len(doc.tables))
